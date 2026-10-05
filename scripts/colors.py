import pandas as pd
import os
from openpyxl import Workbook

# Путь к папке с фотографиями
photo_folder = 'color_images/'

# Путь к файлу Excel
excel_file = 'wtf.xlsx'

# Создание нового Excel-файла
wb = Workbook()
ws = wb.active

# Заголовки столбцов
ws['A1'] = 'Название с фото'
ws['B1'] = 'Название без фото'

# Загрузка данных из Excel
df = pd.read_excel(excel_file)

# Функция для проверки наличия файла фотографии
def check_photo_exists(photo_name):
    photo_path = os.path.join(photo_folder, str(photo_name) + '.jpg')
    return os.path.exists(photo_path)

# Функция для проверки условия в поле DescrizioneArticolo_ItemDescription
def check_description(description):
    return 'BASE' not in description

# Списки для хранения уникальных значений
unique_with_photo = set()
unique_without_photo = set()

# Формирование списка с фотографиями и без
for index, row in df.iterrows():
    description = row['DescrizioneArticolo_ItemDescription']
    if check_description(description):
        photo_name = str(row['CLA3_senza_spazi'])
        if check_photo_exists(photo_name):
            unique_with_photo.add(photo_name)
        else:
            unique_without_photo.add(photo_name)

# Запись данных в Excel
for idx, name in enumerate(unique_with_photo, start=2):
    ws[f'A{idx}'] = name

for idx, name in enumerate(unique_without_photo, start=2):
    ws[f'B{idx}'] = name

# Сохранение Excel-файла
wb.save("colors.xlsx")